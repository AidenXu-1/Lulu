"""C3 封面双入口与用户出口的离屏、零外部触达测试。"""

from __future__ import annotations

from dataclasses import replace
import csv
import hashlib
import os
from pathlib import Path
import socket
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from src.config import AppConfig  # noqa: E402
from src.core import library, obsidian  # noqa: E402
from src.core.collect import export  # noqa: E402
from src.core.collect.contracts import CollectResult  # noqa: E402
from src.core.collect.cover import CoverReceipt  # noqa: E402
from src.core.collect.cover_workflow import CoverWorkflowService  # noqa: E402
from src.core.collect.models import CollectedItem  # noqa: E402
from src.core.collect.preferences import CollectCoverPreferenceService  # noqa: E402
from src.core.collect.repository import CollectionRepository  # noqa: E402
from src.core.models import Transcript  # noqa: E402
from src.core.task import TaskState  # noqa: E402
from src.ui.collection_page import CollectionPage  # noqa: E402
from src.ui.main_window import MainWindow  # noqa: E402


class FakeLocalizer:
    startup_recovery_receipts = ()

    def __init__(self, receipt: CoverReceipt | None = None) -> None:
        self.receipt = receipt or CoverReceipt.failed("cover_timeout")
        self.calls: list[tuple[str, str]] = []

    def localize(self, *, platform, work_id, url, is_cancelled=None):
        self.calls.append((platform, work_id))
        return self.receipt


class BlockingRecoveryLocalizer:
    startup_recovery_receipts = ()

    def __init__(self) -> None:
        self.started = threading.Event()
        self.release = threading.Event()
        self.calls = 0

    def localize(self, **_kwargs):
        self.calls += 1
        self.started.set()
        if not self.release.wait(2):
            raise AssertionError("test did not release recovery localizer")
        return CoverReceipt.failed("cover_commit_uncertain")


class FakeCollectionService:
    def __init__(self) -> None:
        from src.core.collect.application import CollectionApplicationService

        self._pure = CollectionApplicationService(
            adapter_resolver=lambda _platform: (_ for _ in ()).throw(
                AssertionError("adapter must stay unused")
            ),
            session_factory=lambda *_args: (_ for _ in ()).throw(
                AssertionError("browser must stay unused")
            ),
        )
        self.collect_calls = 0

    def inspect(self, raw_input):
        return self._pure.inspect(raw_input)

    def collect(self, *_args, **_kwargs):
        self.collect_calls += 1
        raise AssertionError("cover retry must not recollect")

    def close_active_session(self):
        return True


def cover_item(status: str = "url_only", *, local_path: str = "") -> CollectedItem:
    return CollectedItem(
        "",
        platform="xiaohongshu",
        work_id="fixture-note",
        author_name="公开品牌",
        title="公开作品",
        work_url="https://www.xiaohongshu.com/explore/fixture-note",
        content_kind="image",
        cover_url="https://assets.example.test/cover.jpg",
        cover_status=status,
        cover_local_path=local_path,
        extra_stats={"_field_state_duration_ms": "not_applicable"},
    )


class CoverUserExitTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.old_library_dir = library.LIBRARY_DIR
        library.LIBRARY_DIR = self.root / "library"
        self.external_calls: list[str] = []

        def blocked(name):
            def fail(*_args, **_kwargs):
                self.external_calls.append(name)
                raise AssertionError(f"unexpected external call: {name}")
            return fail

        self.guards = [
            patch.object(socket, "getaddrinfo", blocked("dns")),
            patch.object(socket, "socket", blocked("socket")),
            patch("urllib.request.urlopen", blocked("urlopen")),
            patch("webbrowser.open", blocked("browser")),
        ]
        for guard in self.guards:
            guard.start()

    def tearDown(self):
        for guard in reversed(self.guards):
            guard.stop()
        library.LIBRARY_DIR = self.old_library_dir
        self.temp.cleanup()

    def _wait(self, predicate, timeout: float = 2.0) -> None:
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            self.app.processEvents()
            if predicate():
                self.app.processEvents()
                return
            time.sleep(0.002)
        self.fail("Qt worker did not finish")

    def test_two_entries_share_truth_and_failed_save_rolls_both_back(self):
        cfg = AppConfig()
        cfg.save = lambda: True
        preferences = CollectCoverPreferenceService(cfg)
        localizer = FakeLocalizer()
        main = None
        with (
            patch.object(AppConfig, "load", return_value=cfg),
            patch("src.core.models_manager.app_data_dir", return_value=self.root),
        ):
            main = MainWindow(
                cover_preference_service=preferences,
                cover_localizer=localizer,
                collection_cover_workflow=CoverWorkflowService(localizer),
            )
            page = main._ensure_page_loaded(5)
            self.assertTrue(main.cover_cb.isEnabled())
            self.assertFalse(main.cover_cb.isChecked())
            self.assertFalse(page.cover_cb.isChecked())

            main.cover_cb.setChecked(True)
            self.assertTrue(preferences.value())
            self.assertTrue(page.cover_cb.isChecked())

            cfg.save = lambda: False
            with patch("src.ui.collection_page.dialogs.alert"):
                page.cover_cb.setChecked(False)
            self.assertTrue(preferences.value())
            self.assertTrue(main.cover_cb.isChecked())
            self.assertTrue(page.cover_cb.isChecked())
            self.assertIs(
                page._service._cover_workflow,
                main._collection_cover_workflow,
            )

            retained = self.root / "collection/covers/xiaohongshu/retained.jpg"
            retained.parent.mkdir(parents=True, exist_ok=True)
            retained.write_bytes(b"retained")
            transcript = Transcript(
                text="任务正文",
                segments=[],
                source="fixture",
                meta={
                    "cover_platform": "xiaohongshu",
                    "cover_work_id": "retained",
                    "cover_url": "https://assets.example.test/retained.jpg",
                    "cover_status": "downloaded",
                    "cover_local_path": "xiaohongshu/retained.jpg",
                },
            )
            main._add_task("fixture", "link", "封面引用任务")
            task = main.tasks[-1]
            task.state = TaskState.DONE
            task.transcript = transcript
            main._delete_task(task.id)
            self.assertTrue(retained.exists())

            entry_id = library.save(transcript, "封面引用文稿")
            library.delete(entry_id)
            self.assertTrue(retained.exists())
        if main is not None:
            main.close()
            main.deleteLater()
            self.app.processEvents()
        self.assertEqual(self.external_calls, [])

    def test_collection_cover_retry_updates_master_and_reveal_keeps_file(self):
        managed = self.root / "collection" / "covers"
        relative = "xiaohongshu/fixture.jpg"
        target = managed / relative
        target.parent.mkdir(parents=True)
        target.write_bytes(b"safe-cover")
        receipt = CoverReceipt(
            "downloaded",
            relative_path=relative,
            size=10,
            sha256=hashlib.sha256(b"safe-cover").hexdigest(),
            mime_type="image/jpeg",
        )
        localizer = FakeLocalizer(receipt)
        workflow = CoverWorkflowService(localizer)
        records = CollectionRepository(self.root / "collection" / "items.json")
        service = FakeCollectionService()
        cfg = AppConfig()
        cfg.save = lambda: True
        page = None
        with patch("src.core.models_manager.app_data_dir", return_value=self.root):
            page = CollectionPage(
                cfg,
                collection_service=service,
                records_repository=records,
                cover_preference_service=CollectCoverPreferenceService(cfg),
                cover_workflow=workflow,
            )
            page._set_items([cover_item()])
            page.table.item(0, 0).setCheckState(Qt.CheckState.Checked)
            page._retry_selected_covers()
            self._wait(lambda: page._cover_retry_worker is None)

            self.assertEqual(service.collect_calls, 0)
            self.assertEqual(page._items[0].cover_status, "downloaded")
            self.assertEqual(records.load().items[0].cover_status, "downloaded")
            self.assertTrue(target.exists())

            page.table.item(0, 0).setCheckState(Qt.CheckState.Checked)
            with patch("src.ui.collection_page.subprocess.run") as reveal:
                page._reveal_selected_cover()
                reveal.assert_called_once_with(
                    ["open", "-R", str(target.resolve())],
                    check=False,
                )
            self.assertTrue(target.exists())
        if page is not None:
            page.shutdown_workers()
            page.deleteLater()
            self.app.processEvents()
        self.assertEqual(self.external_calls, [])

    def test_recovery_retry_refreshes_running_copy_and_blocks_duplicate_trigger(self):
        localizer = BlockingRecoveryLocalizer()
        cfg = AppConfig()
        cfg.save = lambda: True
        page = None
        with patch("src.core.models_manager.app_data_dir", return_value=self.root):
            page = CollectionPage(
                cfg,
                collection_service=FakeCollectionService(),
                cover_workflow=CoverWorkflowService(localizer),
            )
            uncertain = replace(
                cover_item("failed"),
                last_error_code="cover:cover_commit_uncertain",
            )
            page._set_items([uncertain])
            page.table.item(0, 0).setCheckState(Qt.CheckState.Checked)
            self.assertEqual(page.cover_retry_btn.text(), "重试恢复")

            page._retry_selected_covers()
            self.assertTrue(localizer.started.wait(1))
            self.app.processEvents()
            self.assertEqual(page.cover_retry_btn.text(), "正在恢复…")
            self.assertFalse(page.cover_retry_btn.isEnabled())

            page._retry_selected_covers()
            self.assertEqual(localizer.calls, 1)

            localizer.release.set()
            self._wait(lambda: page._cover_retry_worker is None)
            self.assertEqual(page.cover_retry_btn.text(), "重试恢复")
            self.assertTrue(page.cover_retry_btn.isEnabled())
        localizer.release.set()
        if page is not None:
            page.shutdown_workers()
            page.deleteLater()
            self.app.processEvents()
        self.assertEqual(self.external_calls, [])


    def test_csv_xlsx_fields_are_local_only_and_bitable_has_no_local_path(self):
        managed = self.root / "collection" / "covers"
        local = managed / "xiaohongshu/export.jpg"
        local.parent.mkdir(parents=True)
        local.write_bytes(b"export")
        item = cover_item("downloaded", local_path="xiaohongshu/export.jpg")
        with patch("src.core.models_manager.app_data_dir", return_value=self.root):
            csv_path = export.write_csv([item], self.root / "items.csv")
            rows = list(csv.reader(csv_path.read_text(encoding="utf-8-sig").splitlines()))
            self.assertEqual(rows[0][-2:], ["封面状态", "封面本地文件"])
            self.assertEqual(rows[1][-2:], ["封面已保存", str(local.resolve())])
            fields = export.to_bitable_records([item])[0]["fields"]
            self.assertNotIn("封面本地文件", fields)
            self.assertNotIn(str(local.resolve()), repr(fields))

            outside = self.root / "outside.jpg"
            outside.write_bytes(b"outside")
            symlink = managed / "xiaohongshu/link.jpg"
            symlink.symlink_to(outside)
            unsafe = cover_item(
                "downloaded",
                local_path="xiaohongshu/link.jpg",
            )
            unsafe_rows = list(csv.reader(export.to_csv_string([unsafe]).splitlines()))
            self.assertEqual(unsafe_rows[1][-2:], ["本地封面文件已丢失", ""])
        self.assertEqual(self.external_calls, [])

    def test_hardlink_is_rejected_by_every_local_exit_while_regular_file_works(self):
        from openpyxl import load_workbook

        platform_root = self.root / "collection" / "covers" / "xiaohongshu"
        platform_root.mkdir(parents=True)
        outside = self.root / "outside-hardlink.jpg"
        outside.write_bytes(b"outside")
        hardlink = platform_root / "hardlink.jpg"
        os.link(outside, hardlink)
        regular = platform_root / "regular.jpg"
        regular.write_bytes(b"regular")
        self.assertEqual(hardlink.stat().st_nlink, 2)
        self.assertEqual(regular.stat().st_nlink, 1)

        hard_item = cover_item(
            "downloaded", local_path="xiaohongshu/hardlink.jpg"
        )
        regular_item = replace(
            cover_item("downloaded", local_path="xiaohongshu/regular.jpg"),
            work_id="regular",
        )
        hard_transcript = Transcript(
            text="硬链接正文",
            segments=[],
            source="fixture",
            meta={
                "cover_platform": "xiaohongshu",
                "cover_work_id": "hardlink",
                "cover_url": "https://assets.example.test/hardlink.jpg",
                "cover_status": "downloaded",
                "cover_local_path": "xiaohongshu/hardlink.jpg",
            },
        )
        regular_transcript = Transcript(
            text="普通正文",
            segments=[],
            source="fixture",
            meta={
                "cover_platform": "xiaohongshu",
                "cover_work_id": "regular",
                "cover_url": "https://assets.example.test/regular.jpg",
                "cover_status": "downloaded",
                "cover_local_path": "xiaohongshu/regular.jpg",
            },
        )
        page = None
        main = None
        with patch("src.core.models_manager.app_data_dir", return_value=self.root):
            real_stat = os.stat
            platform_stat_calls = 0

            def drifting_stat(path, *args, **kwargs):
                nonlocal platform_stat_calls
                result = real_stat(path, *args, **kwargs)
                if (
                    kwargs.get("dir_fd") is None
                    and Path(path) == platform_root
                ):
                    platform_stat_calls += 1
                    if platform_stat_calls == 2:
                        values = list(result)
                        values[1] += 1
                        return os.stat_result(values)
                return result

            with patch.object(export.os, "stat", side_effect=drifting_stat):
                self.assertIsNone(
                    export.resolve_managed_cover_file(
                        "xiaohongshu/regular.jpg"
                    )
                )

            # CSV/XLSX 共用唯一受管路径解析；攻击样本不得输出本地路径。
            csv_rows = list(
                csv.reader(
                    export.to_csv_string([hard_item, regular_item]).splitlines()
                )
            )
            self.assertEqual(csv_rows[1][-2:], ["本地封面文件已丢失", ""])
            self.assertEqual(csv_rows[2][-2:], ["封面已保存", str(regular.resolve())])
            xlsx_path = export.write_xlsx(
                [hard_item, regular_item], self.root / "hardlink.xlsx"
            )
            workbook = load_workbook(xlsx_path, read_only=True)
            try:
                xlsx_rows = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
            self.assertEqual(xlsx_rows[1][-2:], ("本地封面文件已丢失", None))
            self.assertEqual(
                xlsx_rows[2][-2:], ("封面已保存", str(regular.resolve()))
            )

            # Obsidian 对攻击样本只保留原始远端 URL，不暴露硬链接目标。
            hard_note = obsidian.note_body(hard_transcript, "硬链接")
            regular_note = obsidian.note_body(regular_transcript, "普通文件")
            self.assertNotIn(outside.resolve().as_uri(), hard_note)
            self.assertIn("https://assets.example.test/hardlink.jpg", hard_note)
            self.assertIn(regular.resolve().as_uri(), regular_note)

            cfg = AppConfig()
            cfg.save = lambda: True
            page = CollectionPage(
                cfg,
                collection_service=FakeCollectionService(),
            )
            page._set_items([hard_item])
            page.table.item(0, 0).setCheckState(Qt.CheckState.Checked)
            self.assertFalse(page.cover_finder_btn.isEnabled())
            with patch("src.ui.collection_page.subprocess.run") as reveal:
                page._reveal_selected_cover()
                reveal.assert_not_called()

            page._set_items([regular_item])
            page.table.item(0, 0).setCheckState(Qt.CheckState.Checked)
            self.assertTrue(page.cover_finder_btn.isEnabled())
            with patch("src.ui.collection_page.subprocess.run") as reveal:
                page._reveal_selected_cover()
                reveal.assert_called_once_with(
                    ["open", "-R", str(regular.resolve())],
                    check=False,
                )

        if page is not None:
            page.shutdown_workers()
            page.deleteLater()
        if main is not None:
            main.close()
            main.deleteLater()
        self.app.processEvents()
        self.assertEqual(self.external_calls, [])


if __name__ == "__main__":
    unittest.main()
