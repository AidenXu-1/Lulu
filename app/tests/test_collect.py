"""采集模块单测:数据模型解析/派生字段/去重 + CSV 导出。纯本地、无网络。"""

import csv
import io
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from core.collect.models import CollectedItem, from_aweme, dedupe, AWEME_TYPE_IMAGE
from core.collect import export
from core.models import Segment, Transcript


_STATE_NAMES = ("create_time", "duration_ms", "like", "collect", "comment", "share")


def _xhs_item(work_id: str, *, state: str | None = None, image: bool = False) -> CollectedItem:
    extra = {}
    if state:
        extra = {f"_field_state_{name}": state for name in _STATE_NAMES}
    if image:
        extra["_field_state_duration_ms"] = "not_applicable"
    return CollectedItem(
        "",
        platform="xiaohongshu",
        work_id=work_id,
        work_url=f"https://www.xiaohongshu.com/explore/{work_id}",
        content_kind="image" if image else "video",
        extra_stats=extra,
    )


def _video_aweme(aid="111"):
    return {
        "aweme_id": aid, "desc": "团购评价讲解", "create_time": 1781416800,
        "aweme_type": 0,
        "author": {"nickname": "陈雯Momo"},
        "video": {"duration": 467047, "play_addr": {"url_list": ["https://cdn/v.mp4"]}},
        "statistics": {"digg_count": 81, "collect_count": 24, "comment_count": 4, "share_count": 10},
    }


class TestModel(unittest.TestCase):
    def test_from_aweme_maps_nine_fields(self):
        it = from_aweme(_video_aweme())
        self.assertEqual(it.author_name, "陈雯Momo")
        self.assertEqual(it.title, "团购评价讲解")
        self.assertEqual((it.like, it.collect, it.comment, it.share), (81, 24, 4, 10))
        self.assertEqual(it.play_url, "https://cdn/v.mp4")
        self.assertEqual(it.work_url, "https://www.douyin.com/video/111")
        self.assertEqual(it.duration_ms, 467047)

    def test_from_aweme_maps_trusted_cover_for_both_homepage_routes(self):
        aweme = _video_aweme()
        aweme["video"]["cover"] = {
            "url_list": [
                "http://unsafe.example.test/cover.jpg",
                "https://p3.example.test/cover.jpg?signature=public",
            ]
        }
        item = from_aweme(aweme)
        self.assertEqual(
            item.cover_url,
            "https://p3.example.test/cover.jpg?signature=public",
        )
        self.assertEqual(item.cover_status, "url_only")

    def test_from_aweme_rejects_untrusted_or_malformed_cover_candidates(self):
        for candidate in (
            "http://p3.example.test/cover.jpg",
            "https://user:secret@p3.example.test/cover.jpg",
            "https:///cover.jpg",
            "not-a-url",
            123,
        ):
            with self.subTest(candidate=candidate):
                aweme = _video_aweme()
                aweme["video"]["cover"] = {"url_list": [candidate]}
                item = from_aweme(aweme)
                self.assertEqual(item.cover_url, "")
                self.assertEqual(item.cover_status, "not_requested")

    def test_from_aweme_uses_safe_cover_fallback_order(self):
        aweme = _video_aweme()
        aweme["video"].update(
            {
                "cover": {"url_list": ["http://unsafe.test/cover.jpg"]},
                "origin_cover": {
                    "url_list": ["https://p3.example.test/original.jpg"]
                },
                "dynamic_cover": {
                    "url_list": ["https://p3.example.test/dynamic.jpg"]
                },
            }
        )
        item = from_aweme(aweme)
        self.assertEqual(
            item.cover_url,
            "https://p3.example.test/original.jpg",
        )
        self.assertEqual(item.cover_status, "url_only")

    def test_publish_date_and_duration_text(self):
        it = from_aweme(_video_aweme())
        self.assertEqual(it.publish_date, "2026-06-14")   # 北京时间
        self.assertEqual(it.duration_text, "7:47")        # 467047ms ≈ 467s

    def test_image_post_marked_no_audio(self):
        a = _video_aweme("222")
        a["aweme_type"] = AWEME_TYPE_IMAGE
        a["video"]["duration"] = 0
        it = from_aweme(a)
        self.assertTrue(it.is_image)
        self.assertEqual(it.duration_text, "图文")

    def test_missing_fields_default_safely(self):
        it = from_aweme({"aweme_id": "333"})
        self.assertEqual(it.author_name, "")
        self.assertEqual(it.like, 0)
        self.assertEqual(it.publish_date, "")
        self.assertEqual(it.duration_text, "")

    def test_existing_douyin_record_backfills_canonical_work_url(self):
        self.assertEqual(
            CollectedItem("123456").work_url,
            "https://www.douyin.com/video/123456",
        )
        self.assertEqual(CollectedItem("invalid/id").work_url, "")

    def test_dedupe_keeps_order(self):
        items = [CollectedItem("1"), CollectedItem("2"), CollectedItem("1"), CollectedItem("3")]
        ids = [x.aweme_id for x in dedupe(items)]
        self.assertEqual(ids, ["1", "2", "3"])

    def test_field_state_defaults_preserve_old_xhs_and_douyin_meaning(self):
        old_xhs = _xhs_item("legacy")
        old_xhs.like = 7
        old_douyin = CollectedItem("douyin", like=0)
        image = _xhs_item("image", image=True)

        self.assertEqual(old_xhs.field_state("like"), "unknown")
        self.assertIsNone(old_xhs.known_numeric("like"))
        self.assertEqual(old_douyin.field_state("like"), "known")
        self.assertEqual(old_douyin.known_numeric("like"), 0)
        self.assertEqual(image.field_state("duration_ms"), "not_applicable")

    def test_invalid_or_misapplied_explicit_state_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "未知字段状态"):
            CollectedItem("1", extra_stats={"_field_state_like": "maybe"})
        with self.assertRaisesRegex(ValueError, "not_applicable"):
            CollectedItem("1", extra_stats={"_field_state_like": "not_applicable"})
        with self.assertRaisesRegex(ValueError, "已知字段值非法"):
            CollectedItem("1", like=-1, extra_stats={"_field_state_like": "known"})


class TestExport(unittest.TestCase):
    def test_timestamped_transcript_export_uses_readable_start_timecodes(self):
        transcript = Transcript(
            text="第一段 第二段 第三段",
            segments=[
                Segment("第一段", start=0.0, end=4.2),
                Segment("第二段", start=65.9, end=70.0),
                Segment("第三段", start=3661.2, end=3665.0),
            ],
        )
        self.assertEqual(
            export.format_transcript_for_export(transcript),
            "[00:00] 第一段\n[01:05] 第二段\n[01:01:01] 第三段",
        )

    def test_transcript_export_falls_back_to_plain_text_without_timestamps(self):
        transcript = Transcript(
            text="保留当前文稿",
            segments=[Segment("保留当前文稿")],
        )
        self.assertEqual(
            export.format_transcript_for_export(transcript),
            "保留当前文稿",
        )

    def test_export_prefers_cross_platform_stable_key_and_keeps_legacy_douyin_id(self):
        xhs = _xhs_item("shared")
        douyin = CollectedItem("shared")
        transcripts = {
            "xiaohongshu:shared": "小红书文稿",
            "douyin:shared": "抖音稳定键文稿",
            "shared": "旧抖音文稿",
        }
        rows = list(csv.reader(io.StringIO(
            export.to_csv_string([xhs, douyin], transcripts)
        )))
        self.assertEqual(rows[1][10], "小红书文稿")
        self.assertEqual(rows[2][10], "抖音稳定键文稿")

    def test_csv_has_header_and_rows(self):
        csv_text = export.to_csv_string([from_aweme(_video_aweme()), from_aweme(_video_aweme("444"))])
        lines = csv_text.strip().splitlines()
        self.assertEqual(lines[0].split(",")[0], "作者昵称")
        self.assertEqual(len(lines), 3)               # 表头 + 2 行
        self.assertIn("陈雯Momo", lines[1])
        self.assertIn("团购评价讲解", lines[1])

    def test_write_csv_utf8_bom(self):
        import os
        with tempfile.TemporaryDirectory() as d:
            p = export.write_csv([from_aweme(_video_aweme())], os.path.join(d, "out.csv"))
            self.assertTrue(p.exists())
            raw = p.read_bytes()
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))   # BOM,Excel 中文不乱码

    def test_work_and_media_links_are_stable_separate_columns_for_all_platforms(self):
        long_work_url = (
            "https://www.xiaohongshu.com/explore/中文作品"
            "?来源=主页采集&说明=" + "长链接" * 20
        )
        items = [
            CollectedItem(
                "",
                platform="xiaohongshu",
                work_id="x1",
                work_url=long_work_url,
                play_url="https://cdn.example/x1.mp4?token=signed",
                risk_status="rate_limited",
            ),
            CollectedItem(
                "d1",
                platform="douyin",
                work_id="d1",
                work_url="https://www.douyin.com/video/d1",
                play_url="",
            ),
            CollectedItem(
                "",
                platform="bilibili",
                work_id="b1",
                work_url="",
                play_url="https://cdn.example/b1.mp4",
            ),
            CollectedItem(
                "",
                platform="xiaohongshu",
                work_id="x-empty",
                work_url="",
                play_url="",
            ),
        ]
        rows = list(csv.reader(io.StringIO(export.to_csv_string(items))))
        self.assertEqual(rows[0][8:10], ["原作品链接", "音视频直链"])
        self.assertEqual(
            [row[8:10] for row in rows[1:]],
            [
                [long_work_url, "https://cdn.example/x1.mp4?token=signed"],
                ["https://www.douyin.com/video/d1", ""],
                ["", "https://cdn.example/b1.mp4"],
                ["", ""],
            ],
        )

        with tempfile.TemporaryDirectory() as directory:
            path = export.write_xlsx(items, Path(directory) / "links.xlsx")
            from openpyxl import load_workbook

            sheet = load_workbook(path).active
            self.assertEqual(
                [sheet.cell(1, column).value for column in (9, 10)],
                ["原作品链接", "音视频直链"],
            )
            self.assertEqual(
                [
                    [sheet.cell(row, column).value for column in (9, 10)]
                    for row in range(2, 6)
                ],
                [
                    [long_work_url, "https://cdn.example/x1.mp4?token=signed"],
                    ["https://www.douyin.com/video/d1", None],
                    [None, "https://cdn.example/b1.mp4"],
                    [None, None],
                ],
            )

    def test_unknown_known_zero_and_image_are_honest_in_csv_xlsx_and_bitable(self):
        unknown = _xhs_item("unknown")
        known_zero = _xhs_item("known-zero", state="known")
        image = _xhs_item("image", image=True)
        items = [unknown, known_zero, image]

        rows = list(csv.reader(io.StringIO(export.to_csv_string(items))))
        self.assertEqual(rows[1][2:8], ["未获取"] * 6)
        self.assertEqual(rows[2][2:8], ["0"] * 6)
        self.assertEqual(rows[3][2:8], ["未获取", "图文"] + ["未获取"] * 4)

        records = export.to_bitable_records(items)
        unknown_fields = records[0]["fields"]
        for field in ("发布日期", "时长", "点赞", "收藏", "评论", "分享"):
            self.assertNotIn(field, unknown_fields)
        zero_fields = records[1]["fields"]
        self.assertEqual(zero_fields["发布日期"], 0)
        self.assertEqual(zero_fields["时长"], "0")
        for field in ("点赞", "收藏", "评论", "分享"):
            self.assertEqual(zero_fields[field], 0)
        self.assertNotIn("时长", records[2]["fields"])

        with tempfile.TemporaryDirectory() as directory:
            path = export.write_xlsx(items, Path(directory) / "truth.xlsx")
            from openpyxl import load_workbook
            sheet = load_workbook(path).active
            self.assertEqual([sheet.cell(2, col).value for col in range(3, 9)], ["未获取"] * 6)
            self.assertEqual([sheet.cell(3, col).value for col in range(3, 9)], [0] * 6)
            self.assertEqual(sheet.cell(4, 4).value, "图文")


class TestClearLogin(unittest.TestCase):
    """隐私出口:clear_login 必须真的删掉登录档案目录,且对未登录安全。"""

    def _patch_profile(self, path):
        from core.collect import douyin_browser
        self._orig = douyin_browser._profile_dir
        douyin_browser._profile_dir = lambda: path
        self.addCleanup(setattr, douyin_browser, "_profile_dir", self._orig)
        return douyin_browser

    def test_clear_removes_profile_and_login_mark(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            pdir = Path(d) / "douyin_profile"
            pdir.mkdir()
            (pdir / ".logged_in").write_text("ok", encoding="utf-8")
            (pdir / "Cookies").write_text("fake-session", encoding="utf-8")
            db = self._patch_profile(pdir)
            self.assertTrue(db.is_logged_in())
            self.assertTrue(db.clear_login())
            self.assertFalse(pdir.exists())
            self.assertFalse(db.is_logged_in())

    def test_clear_when_not_logged_in_is_safe(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            pdir = Path(d) / "douyin_profile"   # 不创建
            db = self._patch_profile(pdir)
            self.assertFalse(db.clear_login())   # 没东西可删,返回 False,不报错


class TestHumanizedPathB(unittest.TestCase):
    """路B 拟人化/降频/异常检测的纯逻辑(不依赖浏览器)。"""

    def test_url_is_verify(self):
        from core.collect import douyin_browser as db
        self.assertTrue(db._url_is_verify("https://www.douyin.com/verify_center?x=1"))
        self.assertTrue(db._url_is_verify("https://sec.douyin.com/captcha"))
        self.assertFalse(db._url_is_verify("https://www.douyin.com/user/MS4w"))
        self.assertFalse(db._url_is_verify(""))

    def test_fatal_lost_login(self):
        from core.collect import douyin_browser as db
        # 本来有登录,中途 session cookie 没了 → 掉登录(不可恢复,硬停)
        self.assertEqual(db._fatal_reason({"ttwid"}, True, 0), "登录态丢失")

    def test_no_false_logout_when_never_logged_in(self):
        from core.collect import douyin_browser as db
        # 从没登录过(had_session=False),不应误报掉登录
        self.assertIsNone(db._fatal_reason({"ttwid"}, False, 0))

    def test_fatal_api_errors(self):
        from core.collect import douyin_browser as db
        self.assertIsNone(db._fatal_reason({"sessionid"}, True, 1))          # 1 次不停
        self.assertEqual(db._fatal_reason({"sessionid"}, True, 2),           # 连续 2 次 → 疑似风控
                         "接口连续返回异常(疑似风控)")

    def test_fatal_normal_none(self):
        from core.collect import douyin_browser as db
        self.assertIsNone(db._fatal_reason({"sessionid", "ttwid"}, True, 0))

    def test_mode_constants(self):
        from core.collect import douyin_browser as db
        self.assertNotEqual(db.MODE_EFFICIENCY, db.MODE_LOW_RISK)

    def test_pacing_within_ranges(self):
        from core.collect import douyin_browser as db
        for _ in range(200):
            self.assertTrue(db._ROUND_WAIT[0] <= db._human_wait() <= db._ROUND_WAIT[1])
            self.assertTrue(db._MICRO_WAIT[0] <= db._micro_wait() <= db._MICRO_WAIT[1])
            self.assertTrue(db._SCROLL_PX[0] <= db._pick_scroll_amount() <= db._SCROLL_PX[1])

    def test_scroll_step_js_uses_amount(self):
        from core.collect import douyin_browser as db
        js = db._scroll_step_js(777)
        self.assertIn("scrollBy(0, 777)", js)


if __name__ == "__main__":
    unittest.main()
